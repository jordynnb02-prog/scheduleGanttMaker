
import streamlit as st
from openpyxl import load_workbook
from io import BytesIO
import hashlib
import matplotlib.pyplot as plt

from schedule_parser import (
    listAllSO,
    getSOactivities,
    groupSchedule,
    readMasterSchedule
)

from gantt_chart import createGanttChart


# ---------------------------------------------------------
# PAGE SETUP
# ---------------------------------------------------------

st.set_page_config(
    page_title="Project Gantt Generator",
    layout="wide"
)

st.title("Project Gantt Generator")


# ---------------------------------------------------------
# UPLOAD MASTER SCHEDULE
# ---------------------------------------------------------

uploaded_file = st.file_uploader(
    "Upload Master Schedule",
    type=["xlsm", "xlsx"]
)

if uploaded_file is not None:

    # Identify the uploaded file by its contents
    file_bytes = uploaded_file.getvalue()
    file_id = hashlib.sha256(file_bytes).hexdigest()

    # Only process the Excel file if it is new
    if (
        st.session_state.get("file_id") != file_id
        or "fullSchedule" not in st.session_state
        or "so_nums" not in st.session_state
    ):

        with st.spinner("Loading Master Schedule Data..."):

            workbook = load_workbook(
                BytesIO(file_bytes),
                data_only=True
            )

            try:
                if "Tasks" not in workbook.sheetnames:
                    st.error("Tasks worksheet not found.")
                    st.stop()

                sheet = workbook["Tasks"]

                # Process the entire Master Schedule
                fullSchedule = readMasterSchedule(sheet)

                # Extract available SO numbers
                so_nums = listAllSO(sheet)

                # Save processed data for future reruns
                st.session_state["fullSchedule"] = fullSchedule
                st.session_state["so_nums"] = so_nums
                st.session_state["file_id"] = file_id

            finally:
                workbook.close()


    # Reuse the previously processed data
    fullSchedule = st.session_state["fullSchedule"]
    so_nums = st.session_state["so_nums"]

    st.success("Master Schedule ready!")


    # ---------------------------------------------------------
    # SELECT SALES ORDER
    # ---------------------------------------------------------

    selectedSo = st.selectbox(
        "Select Sales Order Number",
        sorted(so_nums),
        index=None,
        placeholder="Search for an SO..."
    )

    if selectedSo is None:
        st.info(
            "Please select a Sales Order to generate a Gantt chart."
        )
        st.stop()


    # ---------------------------------------------------------
    # GET PROJECT ACTIVITIES
    # ---------------------------------------------------------

    projectSchedule = getSOactivities(
        fullSchedule,
        selectedSo
    )



    # ---------------------------------------------------------
    # GROUP ACTIVITIES
    # ---------------------------------------------------------

    grouped = groupSchedule(projectSchedule)

    st.subheader(f"Schedule Sections — {selectedSo}")

    selectedGroups = {}

    for section, tasks in grouped.items():
        if st.checkbox(section, value=True, key=f"{file_id}_{selectedSo}_{section}"):
            selectedGroups[section] = tasks


# ---------------------------------------------------------
# GENERATE GANTT CHART
# ---------------------------------------------------------

    st.subheader("Gantt Chart")
    if selectedGroups:
        with st.spinner("Generating Gantt chart..."):
            fig = createGanttChart(selectedGroups)
            st.pyplot(fig)
            plt.close(fig)
    else:
        st.info("Select at least one section to display the Gantt chart.")
else:
    st.info("Upload a Master Schedule to get started.")
